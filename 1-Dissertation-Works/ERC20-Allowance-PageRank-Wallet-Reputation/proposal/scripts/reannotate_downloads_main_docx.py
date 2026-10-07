# -*- coding: utf-8 -*-
"""Strip PDF-converted comment drawings from Downloads/main.docx and replace
them with proper Word comments (author: Taehong Kwon).

The PDF sticky-note icons and comment textboxes were converted as DrawingML
AlternateContent shapes (small Group icons + Textbox). Those are removed, then
the same examiner-correction comments are re-anchored as native w:comment.
"""
from __future__ import annotations

import re
import shutil
import sys
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

SRC = Path(r"c:\Users\abito\Downloads\main.docx")
BACKUP = SRC.with_name("main.before-word-comments.docx")

AUTHOR = "Taehong Kwon"
INITIALS = "TK"
COMMENT_DATE = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
REL_COMMENTS = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments"
)

# (anchor, occurrence, comment). occurrence: "only" | "first" | "last".
ANNOTATIONS: list[tuple[str, str, str]] = [
    (
        "Table of Contents",
        "only",
        "Examiner comment 2: Table of contents updated for the new Section 3.4 "
        "(Attestation-Based Reputation Tooling in Deployed DeFi Systems).",
    ),
    (
        "Attestation-based reputation: A wallet score obtained by aggregating weighted credentials",
        "only",
        "Examiner comments 1 and 3: Added a definition of attestation-based reputation, so that the way "
        "trust is normally established in a DeFi system (Gitcoin / Human Passport and comparable tooling) "
        "is defined before EndorseRank is introduced.",
    ),
    (
        "Sybil attack: The creation of many illegitimate identities",
        "only",
        "Examiner comment 3: Added a definition of Sybil attack, since Passport-style scoring exists "
        "primarily to address Sybil attacks.",
    ),
    (
        "In current practice, trust in a DeFi setting is established in two broad ways",
        "only",
        "Examiner comments 1 and 2: New paragraph stating how trust is normally established in a DeFi "
        "system: consistent, transparent on-chain activity that is costly to manipulate, and attestation "
        "tooling such as Human Passport (formerly Gitcoin Passport) that combines on-chain and off-chain "
        "data against a passing threshold. It also notes that both mechanisms decide admission rather "
        "than order wallets by financial reliability, and points to the new Section 3.4.",
    ),
    (
        "nor the attestation-based tooling reviewed in Section",
        "only",
        "Examiner comment 2: The novelty claim is no longer an unqualified 'no existing study'. It now "
        "states precisely what the graph-based literature and the deployed attestation tooling each do "
        "not do, so the contribution of EndorseRank is legible against both.",
    ),
    (
        "A third consideration concerns the reputation tooling already deployed",
        "only",
        "Examiner comment 2: The research problem now positions EndorseRank against Passport-style "
        "tooling, distinguishing an admission gate from a graded ordering of wallets.",
    ),
    (
        "Can EndorseRank, a Weighted PageRank computed over ERC-20 allowance",
        "only",
        "Examiner comment 4: Main research question shortened from 65 words to 28. The setting, "
        "perspective, comparison baselines and evaluation criteria were removed from the sentence itself.",
    ),
    (
        "The setting, perspective, comparison baselines, and evaluation criteria",
        "only",
        "Examiner comment 4: The side conditions removed from the research question are carried by the "
        "SPICE table immediately below, so no information is lost.",
    ),
    (
        "Because trust in operational DeFi systems is established not only",
        "only",
        "Examiner comment 2: Section 3.3 now points forward to Section 3.4, so the gap statement is read "
        "against deployed practice and not only against the scholarly literature.",
    ),
    (
        "Attestation-Based Reputation Tooling in Deployed DeFi Systems",
        "only",
        "Examiner comments 1 to 3: New Section 3.4 reviews the reputational tooling actually used in DeFi, "
        "with Gitcoin / Human Passport as the primary case: Stamps as weighted verifiable credentials, a "
        "score compared against a passing threshold, and optional on-chain attestation via the Ethereum "
        "Attestation Service. Sources: Human Passport (n.d.); Bordeianu and Popescu (2026), Applied "
        "Sciences 16(14), 6929; Siddarth et al. (2020), Frontiers in Blockchain 3, 590171; Buterin et al. "
        "(2019), Management Science 65(11).",
    ),
    (
        "Two properties of this design bear directly on the present study",
        "only",
        "Examiner comment 3: States explicitly that Passport considers both on-chain and off-chain data to "
        "calculate a trust score addressing Sybil attacks, and that the result functions as a threshold "
        "gate rather than as a ranking. Credential farming is recorded as its residual weakness.",
    ),
    (
        "EndorseRank is therefore neither a competitor to nor a replacement",
        "only",
        "Examiner comment 2: The original contribution is now stated as an explicit five-way contrast with "
        "Passport-style tooling: signal source, data scope, question answered, output form and validation "
        "target. This is the passage that answers 'how will EndorseRank make a novel contribution'.",
    ),
    (
        "In summary, the literature shows a progression",
        "only",
        "Examiner comment 2: Chapter summary rewritten so that attestation-based tooling sits alongside the "
        "graph-based and machine-learning strands, and the remaining gap is stated across all three.",
    ),
    (
        "No Personhood Layer: EndorseRank is computed entirely from on-chain state",
        "only",
        "Examiner comment 3: Limitation added. EndorseRank uses no off-chain evidence and therefore offers "
        "no personhood guarantee of its own; it is a relational ordering over an already admitted population.",
    ),
    (
        "Untested Composition: Because a personhood gate and a relational ranking",
        "only",
        "Examiner comment 3: Limitation added. Combining a Passport-style personhood gate with EndorseRank "
        "is a plausible deployment pattern but is left to future work, and credential farming would carry "
        "over to that composition.",
    ),
    (
        "A Distinct Contribution Relative to Attestation-Based Tooling",
        "only",
        "Examiner comment 2: The significance chapter now states the original contribution relative to "
        "attestation-based tooling: delegation treated as a measurable, rankable trust relation.",
    ),
    (
        "Bordeianu, A. A., & Popescu, D. E. (2026)",
        "only",
        "Examiner comments 2 and 3: Added Bordeianu and Popescu (2026), Applied Sciences 16(14), 6929, a "
        "peer-reviewed comparative analysis that covers Gitcoin's Human Passport. DOI verified against "
        "Crossref: 10.3390/app16146929.",
    ),
    (
        "Buterin, V., Hitzig, Z., & Weyl, E. G. (2019)",
        "only",
        "Examiner comment 2: Added Buterin, Hitzig and Weyl (2019), Management Science 65(11), 5171-5187, "
        "for the quadratic-funding setting that Passport was built to protect. DOI verified against "
        "Crossref: 10.1287/mnsc.2019.3337.",
    ),
    (
        "Human Passport developer documentation",
        "only",
        "Examiner comments 1 to 3: Added the Human Passport (formerly Gitcoin Passport) developer "
        "documentation as the primary source for the Stamps, weighted score and threshold mechanism.",
    ),
    (
        "Siddarth, D., Ivliev, S., Siri, S., & Berman, P. (2020)",
        "only",
        "Examiner comment 3: Added Siddarth et al. (2020), Frontiers in Blockchain 3, 590171, a "
        "peer-reviewed review of Sybil-resistant proof-of-personhood protocols. DOI verified against "
        "Crossref: 10.3389/fbloc.2020.590171.",
    ),
]


def norm(text: str) -> str:
    return re.sub(r"[^a-z0-9]", "", text.lower())


def is_comment_drawing(ac) -> bool:
    """True if this AlternateContent is a PDF sticky-note icon or comment textbox."""
    texts = "".join(t.text or "" for t in ac.findall(".//{*}t"))
    if "Taehong Kwon" in texts or "Examiner comment" in texts:
        return True

    drawing = ac.find(".//{*}drawing")
    if drawing is None:
        return False
    ext = drawing.find(".//{*}ext")
    if ext is None:
        return False
    try:
        cx, cy = int(ext.get("cx") or 0), int(ext.get("cy") or 0)
    except ValueError:
        return False

    for cnv in drawing.findall(".//{*}cNvPr"):
        name = cnv.get("name") or ""
        # sticky-note cursor icon ~0.2in square
        if name.startswith("Group") and cx <= 250000 and cy <= 250000:
            return True
        # comment popup textbox ~2.54in x 1.27in
        if name.startswith("Textbox") and 2000000 <= cx <= 3000000 and 1000000 <= cy <= 1500000:
            return True
    return False


def strip_comment_drawings(docx_path: Path) -> tuple[int, int]:
    """Remove comment icon/textbox drawings. Returns (removed_ac, removed_empty_runs)."""
    tmp = docx_path.with_suffix(".strip_tmp.docx")
    removed_ac = 0
    removed_runs = 0

    with zipfile.ZipFile(docx_path, "r") as zin, zipfile.ZipFile(
        tmp, "w", compression=zipfile.ZIP_DEFLATED
    ) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename != "word/document.xml":
                zout.writestr(item, data)
                continue

            root = etree.fromstring(data)
            for ac in list(root.findall(".//{*}AlternateContent")):
                if not is_comment_drawing(ac):
                    continue
                parent = ac.getparent()
                if parent is None:
                    continue
                parent.remove(ac)
                removed_ac += 1

            # Drop runs left with only rPr after drawing removal.
            for run in list(root.findall(".//{*}r")):
                kids = [c for c in run if etree.QName(c).localname != "rPr"]
                if kids:
                    continue
                parent = run.getparent()
                if parent is not None and etree.QName(parent).localname == "p":
                    parent.remove(run)
                    removed_runs += 1

            data = etree.tostring(
                root, xml_declaration=True, encoding="UTF-8", standalone=True
            )
            zout.writestr(item, data)

    shutil.move(str(tmp), str(docx_path))
    return removed_ac, removed_runs


def ensure_comments_part(docx_path: Path) -> None:
    tmp = docx_path.with_suffix(".comments_tmp.docx")
    with zipfile.ZipFile(docx_path, "r") as zin, zipfile.ZipFile(
        tmp, "w", compression=zipfile.ZIP_DEFLATED
    ) as zout:
        names = set(zin.namelist())
        ct = etree.fromstring(zin.read("[Content_Types].xml"))
        rels = etree.fromstring(zin.read("word/_rels/document.xml.rels"))

        if not any(el.get("PartName") == "/word/comments.xml" for el in ct):
            etree.SubElement(
                ct,
                f"{{{CT_NS}}}Override",
                PartName="/word/comments.xml",
                ContentType=(
                    "application/vnd.openxmlformats-officedocument"
                    ".wordprocessingml.comments+xml"
                ),
            )

        if not any(el.get("Type") == REL_COMMENTS for el in rels):
            used = {el.get("Id") for el in rels}
            rid = "rIdComments"
            n = 90
            while rid in used:
                rid = f"rId{n}"
                n += 1
            etree.SubElement(
                rels,
                f"{{{REL_NS}}}Relationship",
                Id=rid,
                Type=REL_COMMENTS,
                Target="comments.xml",
            )

        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = etree.tostring(
                    ct, xml_declaration=True, encoding="UTF-8", standalone=True
                )
            elif item.filename == "word/_rels/document.xml.rels":
                data = etree.tostring(
                    rels, xml_declaration=True, encoding="UTF-8", standalone=True
                )
            elif item.filename == "word/comments.xml":
                continue
            zout.writestr(item, data)

        if "word/comments.xml" not in names:
            comments = etree.Element(f"{{{W_NS}}}comments", nsmap={"w": W_NS})
            zout.writestr(
                "word/comments.xml",
                etree.tostring(
                    comments, xml_declaration=True, encoding="UTF-8", standalone=True
                ),
            )

    shutil.move(str(tmp), str(docx_path))


def add_comment_element(comments_root, cid: int, text: str) -> None:
    c = OxmlElement("w:comment")
    c.set(qn("w:id"), str(cid))
    c.set(qn("w:author"), AUTHOR)
    c.set(qn("w:date"), COMMENT_DATE)
    c.set(qn("w:initials"), INITIALS)
    p = OxmlElement("w:p")
    ref_run = OxmlElement("w:r")
    ref_run.append(OxmlElement("w:annotationRef"))
    p.append(ref_run)
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = text
    run.append(t)
    p.append(run)
    c.append(p)
    comments_root.append(c)


def wrap_paragraph_comment(paragraph, cid: int) -> None:
    p = paragraph._p
    start = OxmlElement("w:commentRangeStart")
    start.set(qn("w:id"), str(cid))
    end = OxmlElement("w:commentRangeEnd")
    end.set(qn("w:id"), str(cid))
    ref_run = OxmlElement("w:r")
    ref = OxmlElement("w:commentReference")
    ref.set(qn("w:id"), str(cid))
    ref_run.append(ref)

    pPr = p.find(qn("w:pPr"))
    if pPr is not None:
        pPr.addnext(start)
    else:
        p.insert(0, start)
    p.append(end)
    p.append(ref_run)


def plain_paragraph_text(paragraph) -> str:
    parts = []
    for node in paragraph._p.iter():
        if etree.QName(node).localname != "t":
            continue
        ancestors = {etree.QName(a).localname for a in node.iterancestors()}
        if ancestors & {"drawing", "txbxContent", "pict", "wsp", "shape"}:
            continue
        if node.text:
            parts.append(node.text)
    return "".join(parts)


def resolve(paragraphs, anchor: str, occurrence: str) -> int:
    needle = norm(anchor)
    hits = [
        i
        for i, p in enumerate(paragraphs)
        if needle in norm(plain_paragraph_text(p))
    ]
    if not hits:
        raise SystemExit(f"Anchor not found: {anchor!r}")
    if occurrence == "first":
        return hits[0]
    if occurrence == "last":
        return hits[-1]
    if len(hits) == 1:
        return hits[0]
    exact = [i for i in hits if norm(plain_paragraph_text(paragraphs[i])) == needle]
    if len(exact) == 1:
        return exact[0]
    raise SystemExit(f"Anchor ambiguous ({hits}): {anchor!r}")


def main() -> None:
    if not SRC.is_file():
        raise SystemExit(f"Missing {SRC}")

    try:
        shutil.copy2(SRC, BACKUP)
    except PermissionError:
        raise SystemExit(f"Cannot write backup {BACKUP}. Close Word and retry.")

    print(f"Backup: {BACKUP}")
    removed_ac, removed_runs = strip_comment_drawings(SRC)
    print(f"Removed {removed_ac} comment drawing(s), {removed_runs} empty run(s)")

    root = etree.fromstring(zipfile.ZipFile(SRC).read("word/document.xml"))
    leftover = [ac for ac in root.findall(".//{*}AlternateContent") if is_comment_drawing(ac)]
    if leftover:
        raise SystemExit(f"{len(leftover)} comment drawing(s) still present")

    ensure_comments_part(SRC)
    doc = Document(str(SRC))

    comments_part = None
    for rel in doc.part.rels.values():
        if rel.reltype == REL_COMMENTS:
            comments_part = rel.target_part
            break
    if comments_part is None:
        raise SystemExit("comments part missing after ensure_comments_part")

    comments_root = comments_part.element
    for old in list(comments_root.findall(qn("w:comment"))):
        comments_root.remove(old)

    cid = 0
    for anchor, occurrence, comment in ANNOTATIONS:
        idx = resolve(doc.paragraphs, anchor, occurrence)
        add_comment_element(comments_root, cid, comment)
        wrap_paragraph_comment(doc.paragraphs[idx], cid)
        print(f"  [{cid}] p{idx}: {anchor[:60]}")
        cid += 1

    try:
        doc.save(str(SRC))
    except PermissionError:
        raise SystemExit(
            f"{SRC} is open in Word. Close it and re-run. "
            f"Restore from {BACKUP} if the file was left half-updated."
        )

    with zipfile.ZipFile(SRC) as z:
        if z.testzip() is not None:
            raise SystemExit("DOCX zip integrity check failed")
        body = z.read("word/document.xml").decode("utf-8")
        comments = etree.fromstring(z.read("word/comments.xml"))
        found = comments.findall(qn("w:comment"))
        authors = {c.get(qn("w:author")) for c in found}
        root = etree.fromstring(body.encode("utf-8"))
        leftover = [
            ac for ac in root.findall(".//{*}AlternateContent") if is_comment_drawing(ac)
        ]
        starts = len(root.findall(".//{*}commentRangeStart"))
        refs = len(root.findall(".//{*}commentReference"))
        remaining_drawings = len(root.findall(".//{*}drawing"))

    if leftover:
        raise SystemExit(f"Comment drawings still present: {len(leftover)}")
    if authors != {AUTHOR}:
        raise SystemExit(f"Unexpected authors: {authors}")
    if len(found) != len(ANNOTATIONS):
        raise SystemExit(f"Expected {len(ANNOTATIONS)} comments, found {len(found)}")
    if starts != refs or starts != len(ANNOTATIONS):
        raise SystemExit(f"Anchor mismatch start={starts} ref={refs}")
    if "ns0:" in body:
        raise SystemExit("document.xml contains ns0: prefix")

    print(
        f"Done. {len(found)} Word comments by {AUTHOR}; "
        f"{remaining_drawings} non-comment drawing(s) kept; wrote {SRC}"
    )


if __name__ == "__main__":
    sys.exit(main())
