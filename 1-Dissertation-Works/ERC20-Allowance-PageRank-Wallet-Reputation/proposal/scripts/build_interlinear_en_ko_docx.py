# -*- coding: utf-8 -*-
"""Build interlinear EN+KO DOCX: each English paragraph followed by matched Korean.

Matching is by exact English text → translation map (never by sequential index).
Supports body paragraphs and table cell paragraphs.
"""
from __future__ import annotations

import argparse
import json
import re
import shutil
import sys
import time
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from deep_translator import GoogleTranslator
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from lxml import etree

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

PROTECT = [
    "Adaptive Weighted PageRank",
    "EndorseRank",
    "PageRank",
    "RiskProp",
    "Arbitrum One",
    "BigQuery",
    "PositionDecrease",
    "PositionIncrease",
    "Human Passport",
    "Gitcoin Passport",
    "basePnlUsd",
    "orderType",
    "ERC-20",
    "EIP-2612",
    "GMX V2",
    "GMX",
    "AWP",
    "DeFi",
    "Sybil",
    "SPICE",
    "UNISA",
    "CSET",
    "Kendall",
    "Spearman",
    "Taehong Kwon",
    "Ernest Mnkandla",
    "Donatien Koulla Moulla",
    "David Sena Attipoe",
]

MANUAL = {
    "by": "지은이",
    "TAEHONG KWON": "권태홍 (TAEHONG KWON)",
    "submitted in accordance with the requirements for the degree of": "다음 학위 요건에 따라 제출함",
    "DOCTOR OF PHILOSOPHY": "철학박사 (Doctor of Philosophy)",
    "in the subject": "전공",
    "COMPUTER SCIENCE": "컴퓨터과학 (Computer Science)",
    "at the": "（기관）",
    "UNIVERSITY OF SOUTH AFRICA": "남아프리카공화국 대학교 (University of South Africa)",
    "SUPERVISOR: Prof. Ernest Mnkandla": "지도교수: Ernest Mnkandla 교수",
    "CO-SUPERVISOR: Prof. Donatien Koulla Moulla": "공동지도교수: Donatien Koulla Moulla 교수",
    "Abstract": "초록",
    "Table of Contents": "목차",
    "Introduction and background": "서론 및 배경",
    "Research Problem": "연구 문제",
    "Research Questions": "연구 질문",
    "Objectives": "연구 목표",
    "Specifications:": "명세:",
    "Chapter 1": "제1장",
    "Chapter 2": "제2장",
    "Chapter 3": "제3장",
    "Chapter 4": "제4장",
    "Chapter 5": "제5장",
    "Chapter 6": "제6장",
    "Chapter 7": "제7장",
    "Chapter 8": "제8장",
    "References": "참고문헌",
    "Appendix": "부록",
    "Appendices": "부록",
    "INTEGRATING ERC-20 ALLOWANCE EDGES INTO PAGERANK FOR ENHANCED ON-CHAIN WALLET REPUTATION SCORING": (
        "향상된 온체인 지갑 평판 산정을 위한 ERC-20 승인(allowance) 간선과 PageRank의 통합"
    ),
}


def load_cache(path: Path) -> dict[str, str]:
    if path.exists():
        return json.loads(path.read_text(encoding="utf-8"))
    return {}


def save_cache(path: Path, cache: dict[str, str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(cache, ensure_ascii=False, indent=1), encoding="utf-8")


def should_skip(text: str) -> bool:
    t = text.strip()
    if not t:
        return True
    if re.fullmatch(r"[\d\W_]+", t, flags=re.UNICODE):
        return True
    # no Latin/Hangul letters → keep as-is (math fragments, symbols)
    if not re.search(r"[A-Za-z\uac00-\ud7a3]", t):
        return True
    hangul = len(re.findall(r"[\uac00-\ud7a3]", t))
    latin = len(re.findall(r"[A-Za-z]", t))
    if hangul > 0 and hangul >= max(latin, 1):
        return True
    return False


def keep_as_is_fragment(text: str) -> bool:
    """Tab/cell fragments that should not be sent to the translator."""
    t = (text or "").strip()
    if not t:
        return True
    if len(t) <= 2 and not re.search(r"[A-Za-z\uac00-\ud7a3]", t):
        return True
    if re.fullmatch(r"[\d\W_∼≈≒≤≥±×÷]+", t, flags=re.UNICODE):
        return True
    if should_skip(text):
        return True
    return False


def protect(text: str) -> tuple[str, list[str]]:
    held: list[str] = []
    out = text
    for term in PROTECT:
        while term in out:
            tok = f"⟦T{len(held)}⟧"
            held.append(term)
            out = out.replace(term, tok, 1)
    return out, held


def unprotect(text: str, held: list[str]) -> str:
    out = text
    for i, val in enumerate(held):
        for tok in (f"⟦T{i}⟧", f"[T{i}]", f"(T{i})"):
            out = out.replace(tok, val)
    return out


def translate_fresh(translator: GoogleTranslator, text: str) -> str:
    protected, held = protect(text)
    max_len = 4500
    if "\t" in protected:
        parts = protected.split("\t")
        out_parts = []
        for part in parts:
            if keep_as_is_fragment(part):
                out_parts.append(part)
            else:
                out_parts.append(_translate_chunk(translator, part, max_len))
        return unprotect("\t".join(out_parts), held)
    if keep_as_is_fragment(protected):
        return unprotect(protected, held)
    return unprotect(_translate_chunk(translator, protected, max_len), held)


def _translate_chunk(translator: GoogleTranslator, text: str, max_len: int) -> str:
    if len(text) <= max_len:
        parts = [text]
    else:
        parts = []
        buf = ""
        for sent in re.split(r"(?<=[.!?])\s+", text):
            if len(buf) + len(sent) + 1 <= max_len:
                buf = f"{buf} {sent}".strip()
            else:
                if buf:
                    parts.append(buf)
                buf = sent
        if buf:
            parts.append(buf)
    out: list[str] = []
    for part in parts:
        translated: str | None = None
        for attempt in range(8):
            try:
                translated = translator.translate(part)
                if translated is None or not str(translated).strip():
                    raise RuntimeError("empty translation result")
                break
            except Exception as exc:  # noqa: BLE001
                wait = 2.0 * (attempt + 1)
                print(f"  retry {wait:.1f}s: {exc}", flush=True)
                time.sleep(wait)
        if translated is None or not str(translated).strip():
            raise RuntimeError(f"translate failed: {part[:80]!r}")
        out.append(str(translated))
        time.sleep(0.15)
    return " ".join(out)


def resolve_ko(en: str, cache: dict[str, str], translator: GoogleTranslator) -> str | None:
    stripped = en.strip()
    if stripped in MANUAL:
        return MANUAL[stripped]
    if en in MANUAL:
        return MANUAL[en]
    if en in cache:
        return cache[en]
    if should_skip(en) or keep_as_is_fragment(en):
        return None
    try:
        ko = translate_fresh(translator, en)
    except Exception as exc:  # noqa: BLE001
        print(f"  skip failed string ({exc}): {en[:60]!r}", flush=True)
        cache[en] = en  # mark attempted; avoid re-hitting API
        return None
    if not ko or not str(ko).strip():
        cache[en] = en
        return None
    cache[en] = ko
    return ko


def insert_paragraph_after(paragraph: Paragraph, text: str, *, italic: bool = True) -> Paragraph:
    new_p = OxmlElement("w:p")
    pPr = paragraph._p.find(qn("w:pPr"))
    if pPr is not None:
        new_p.insert(0, etree.fromstring(etree.tostring(pPr)))
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    run = new_para.add_run(text)
    run.italic = italic
    return new_para


def iter_body_and_table_paragraphs(doc: Document):
    yield from doc.paragraphs
    for table in doc.tables:
        yield from iter_table_paragraphs(table)


def iter_table_paragraphs(table: Table):
    for row in table.rows:
        for cell in row.cells:
            yield from cell.paragraphs
            for nested in cell.tables:
                yield from iter_table_paragraphs(nested)


def append_title_comment(dst: Path, note: str) -> None:
    with zipfile.ZipFile(dst, "r") as zin:
        blob = {n: zin.read(n) for n in zin.namelist()}
    if "word/comments.xml" not in blob:
        # create minimal comments part is complex; skip if absent
        comments_xml = (
            b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            b'<w:comments xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"></w:comments>'
        )
        # Without relationship wiring Word may ignore; only append if comments exist
        return

    comments = etree.fromstring(blob["word/comments.xml"])
    ids = [int(c.get(f"{W}id")) for c in comments.findall(f"{W}comment") if c.get(f"{W}id")]
    new_id = str(max(ids) + 1 if ids else 0)
    now = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    comment = OxmlElement("w:comment")
    comment.set(qn("w:id"), new_id)
    comment.set(qn("w:author"), "Taehong Kwon")
    comment.set(qn("w:initials"), "TK")
    comment.set(qn("w:date"), now)
    p = OxmlElement("w:p")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = note
    r.append(t)
    p.append(r)
    comment.append(p)
    comments.append(comment)
    blob["word/comments.xml"] = etree.tostring(comments, xml_declaration=True, encoding="UTF-8", standalone=True)

    doc = etree.fromstring(blob["word/document.xml"])
    body = doc.find(f"{W}body")
    target = None
    for child in body:
        if etree.QName(child).localname != "p":
            continue
        if child.find(f".//{W}t") is not None:
            target = child
            break
    if target is not None:
        cstart = OxmlElement("w:commentRangeStart")
        cstart.set(qn("w:id"), new_id)
        cend = OxmlElement("w:commentRangeEnd")
        cend.set(qn("w:id"), new_id)
        cref_r = OxmlElement("w:r")
        cref = OxmlElement("w:commentReference")
        cref.set(qn("w:id"), new_id)
        cref_r.append(cref)
        target.insert(0, cstart)
        target.append(cend)
        target.append(cref_r)
        blob["word/document.xml"] = etree.tostring(doc, xml_declaration=True, encoding="UTF-8", standalone=True)

    tmp = dst.with_suffix(".tmp.zip")
    with zipfile.ZipFile(tmp, "w", compression=zipfile.ZIP_DEFLATED) as zout:
        for n, data in blob.items():
            zout.writestr(n, data)
    tmp.replace(dst)


def pretranslate_unique(
    texts: list[str],
    cache: dict[str, str],
    translator: GoogleTranslator,
    cache_path: Path,
) -> None:
    unique = []
    seen = set()
    for t in texts:
        if not t.strip() or t in seen:
            continue
        if t.strip() in MANUAL or t in MANUAL or t in cache:
            continue
        if should_skip(t) or keep_as_is_fragment(t):
            continue
        seen.add(t)
        unique.append(t)
    print(f"Unique strings to translate: {len(unique)} (cache size {len(cache)})", flush=True)
    for i, t in enumerate(unique, 1):
        resolve_ko(t, cache, translator)
        if i % 25 == 0 or i == len(unique):
            save_cache(cache_path, cache)
            print(f"  translated {i}/{len(unique)}", flush=True)


def build(src: Path, dst: Path, cache_path: Path, note: str) -> int:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    if not src.exists():
        print("Missing source", src)
        return 1
    if dst.exists():
        dst.unlink()

    # Merge proposal cache hits when available
    cache = load_cache(cache_path)
    prop_cache = Path(__file__).resolve().parent / "_ko_translation_cache.json"
    if prop_cache.exists():
        for k, v in load_cache(prop_cache).items():
            cache.setdefault(k, v)

    shutil.copy2(src, dst)
    translator = GoogleTranslator(source="en", target="ko")
    doc = Document(str(dst))

    # Snapshot paragraphs before mutation
    paras = list(iter_body_and_table_paragraphs(doc))
    en_texts = [p.text for p in paras]
    pretranslate_unique(en_texts, cache, translator, cache_path)

    pairs: list[tuple[Paragraph, str]] = []
    for p in paras:
        en = p.text
        if not en.strip():
            continue
        ko = resolve_ko(en, cache, translator)
        if not ko or ko.strip() == en.strip():
            continue
        pairs.append((p, ko))

    for p, ko in reversed(pairs):
        insert_paragraph_after(p, ko, italic=True)

    save_cache(cache_path, cache)
    doc.save(str(dst))
    append_title_comment(dst, note)

    with zipfile.ZipFile(dst, "r") as z:
        bad = z.testzip()
        print("testzip:", bad)
        if bad is not None:
            print("WARNING: testzip reported", bad, "- rechecking after reopen")
            # python-docx save can leave transient CRC issues; reopen validate below

    with zipfile.ZipFile(dst, "r") as z:
        bad2 = z.testzip()
        print("testzip recheck:", bad2)
        if bad2 is not None:
            raise RuntimeError(f"corrupt docx: {bad2}")
        if b"ns0:" in z.read("word/document.xml"):
            raise RuntimeError("ns0 prefix found in document.xml")

    # Adjacent match verification on body paragraphs only
    d2 = Document(str(dst))
    texts = [p.text for p in d2.paragraphs]
    ok = 0
    badn = 0
    i = 0
    while i < len(texts) - 1:
        en = texts[i]
        if not en.strip():
            i += 1
            continue
        expected = resolve_ko(en, cache, translator)
        nxt = texts[i + 1]
        if expected and nxt == expected:
            ok += 1
            i += 2
            continue
        i += 1
    print(f"Inserted KO under {len(pairs)} English blocks")
    print(f"Body adjacent exact matches: {ok}")
    print("DONE:", dst)
    print("bytes:", dst.stat().st_size)
    return 0


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--src", required=True)
    ap.add_argument("--dst", required=True)
    ap.add_argument("--cache", required=True)
    ap.add_argument("--note", default="EN-KO interlinear translation (Taehong Kwon).")
    args = ap.parse_args()
    return build(Path(args.src), Path(args.dst), Path(args.cache), args.note)


if __name__ == "__main__":
    raise SystemExit(main())
