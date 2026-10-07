"""Embed Figure 1 (conceptual model) from proposal PDF into the pandoc DOCX export."""

from __future__ import annotations

import sys
import zipfile
from io import BytesIO
from pathlib import Path

import fitz
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches


FIGURE_CAPTION = "Figure 1: EndorseRank vs. AWP: hypotheses and evaluation constructs"


def extract_figure_png(pdf_path: Path) -> bytes:
    doc = fitz.open(pdf_path)
    page = None
    for candidate in doc:
        if FIGURE_CAPTION in candidate.get_text():
            page = candidate
            break
    if page is None:
        raise RuntimeError("Figure 1 caption not found in PDF")

    clip = fitz.Rect(72, 115, page.rect.width - 72, 285)
    pix = page.get_pixmap(clip=clip, dpi=200)
    return pix.tobytes("png")


def embed_figure(docx_path: Path, png_bytes: bytes) -> None:
    document = Document(docx_path)
    target = None
    for paragraph in document.paragraphs:
        if paragraph.text.strip() == FIGURE_CAPTION:
            target = paragraph
            break
    if target is None:
        raise RuntimeError(f"Caption paragraph not found in DOCX: {FIGURE_CAPTION!r}")

    image_stream = BytesIO(png_bytes)
    image_paragraph = target.insert_paragraph_before()
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = image_paragraph.add_run()
    run.add_picture(image_stream, width=Inches(6.0))

    document.save(docx_path)


def verify_docx(docx_path: Path) -> None:
    with zipfile.ZipFile(docx_path) as zf:
        if zf.testzip() is not None:
            raise RuntimeError("DOCX zip integrity check failed")
        names = zf.namelist()
        if not any(name.startswith("word/media/") for name in names):
            raise RuntimeError("No embedded media found after figure insert")


def main() -> None:
    if len(sys.argv) != 4:
        raise SystemExit("Usage: embed_figure_in_docx.py <main.pdf> <input.docx> <output.docx>")

    pdf_path = Path(sys.argv[1])
    input_docx = Path(sys.argv[2])
    output_docx = Path(sys.argv[3])

    png_bytes = extract_figure_png(pdf_path)
    if input_docx.resolve() != output_docx.resolve():
        output_docx.write_bytes(input_docx.read_bytes())
    embed_figure(output_docx, png_bytes)
    verify_docx(output_docx)


if __name__ == "__main__":
    main()
