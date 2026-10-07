# -*- coding: utf-8 -*-
"""Fill student sections of UNISA Supervision Agreement; author comments as Taehong Kwon."""
from __future__ import annotations

import shutil
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

SRC = Path(r"c:\Users\abito\Downloads\Supervisor_Agreement_template[40].docx")
OUT = SRC
BACKUP = SRC.with_name(SRC.stem + "_blank_backup.docx")
REPO_COPY = Path(
    r"D:\Github\phd_works\2-Graduation\official-letters"
    r"\2026-07-28-supervisor-agreement-student-completed.docx"
)

AUTHOR = "Taehong Kwon"
INITIALS = "TK"
COMMENT_DATE = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
REL_COMMENTS = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments"
)

TITLE = (
    "Investigating the predictive power of on-chain reputation for debt repayment "
    "in DeFi: A case study of Aave's credit delegation"
)
DESCRIPTION = (
    "This doctoral study develops and evaluates EndorseRank, an on-chain wallet "
    "reputation score obtained by applying PageRank to a directed endorsement graph "
    "built from latest ERC-20 allowance edges on Arbitrum One. EndorseRank is compared "
    "with Adaptive Weighted PageRank (AWP) and related baselines using rank-correlation "
    "alignment (Spearman rho, Kendall tau) across proxy families (allowance, transfer, "
    "risk-related, and trading-success signals including GMX V2). The work addresses "
    "scalability, interpretability, and domain alignment of social reputation methods "
    "for DeFi settings."
)
B41 = (
    "I confirm the supervision and communication arrangements set out above. I am a "
    "distance student based in the Republic of Korea and will communicate primarily by "
    "e-mail and online meetings, arranging face-to-face meetings when feasible (at least "
    "once a year if required). I expect guidance on research design, methodology, and "
    "thesis structure; coordinated feedback from the supervisor and co-supervisors within "
    "the stated time frames; and support toward proposal/thesis milestones, ethics "
    "clearance where required, manuscript submissions, and examination. I have personal "
    "access to computing facilities and the Internet as expected for School of Computing "
    "students."
)
B51 = (
    "Phase 1 (proposal / foundation): Finalise the research proposal and any required "
    "ethics clearance; consolidate literature and methodology. "
    "Phase 2 (empirical work / drafting): Complete empirical evaluation on Arbitrum One "
    "cohorts; draft thesis chapters; prepare two manuscripts for submission to accredited "
    "journals. "
    "Phase 3 (completion): Thesis revision, professional language editing, examination "
    "submission, and viva voce if applicable to this cohort. "
    "Time commitment: approximately 20-30 hours per week for research and writing, with "
    "draft chapters submitted according to agreed milestones and revised versions returned "
    "within six weeks where applicable."
)
A3 = (
    "Doctoral researcher in Computer Science focusing on on-chain reputation scoring "
    "(EndorseRank) for DeFi. Experience includes designing and running large-scale "
    "blockchain data pipelines (Arbitrum One / BigQuery), implementing PageRank-based "
    "reputation methods, and writing a full dissertation draft under Prof. Ernest Mnkandla "
    "with co-supervision by Dr Donatien Koulla Moulla; proposal drafts have also received "
    "detailed feedback from Dr David Sena Attipoe."
)
A2 = (
    "Currently registered for the Doctor of Philosophy in Computer Science (98803) at "
    "the University of South Africa."
)

FILLS: dict[int, tuple[str, str]] = {
    9: (
        "(Name of graduate student) Taehong Kwon",
        "학생란: 성명 Taehong Kwon 기입",
    ),
    11: (
        "(Signature) Taehong Kwon",
        "학생란: 서명란에 성명 기입(필요 시 자필/전자서명으로 교체)",
    ),
    13: (
        "(Date) 28 July 2026",
        "학생란: 서명 일자 2026-07-28",
    ),
    15: (
        "(Name of supervisor) Prof. Ernest Mnkandla",
        "지도교수명 기입(서명란은 교수용으로 비움). 등록 서한 기준.",
    ),
    26: (
        "A1 Full name of candidate: Taehong Kwon",
        "A1: 성명 기입",
    ),
    27: (
        "Surname: Kwon",
        "A1: Surname",
    ),
    28: (
        "First names: Taehong",
        "A1: First names",
    ),
    30: (
        "A2 Academic and professional qualifications:",
        "A2: 제목 행 유지",
    ),
    31: (
        A2,
        "A2: 레포에 학사/석사 원문 없음 → 현재 UNISA PhD 등록만 기입. 이전 학위 있으면 추가 요망.",
    ),
    33: (
        "A3 Candidate's experience:",
        "A3: 제목 행 유지",
    ),
    34: (
        A3,
        "A3: 연구 경험 요약(레포·지도체제 기준)",
    ),
    36: (
        "Title of topic: " + TITLE,
        "A4: 등록 수리 서한 working title 사용(현재 논문 표지 제목과 다를 수 있음)",
    ),
    38: (
        "Brief project description: " + DESCRIPTION,
        "A4: 논문 Abstract/연구 요지 기반 프로젝트 설명",
    ),
    40: (
        "Module TFCOS01",
        "A4: 등록 과목 코드 TFCOS01",
    ),
    44: (
        "A5 Personal particulars: Taehong Kwon",
        "A5: 후보 성명",
    ),
    45: (
        "Student number: 28576810",
        "A5: 학번(등록 서한 2857-681-0)",
    ),
    46: (
        "Degree registered for: PhD (Computer Science) (98803)",
        "A5: 학위·자격코드",
    ),
    47: (
        "Postal address: 19, Munhwawon-ro 38beon-gil, Room 206, Yuseong-gu, "
        "Daejeon 34167, Republic of Korea",
        "A5: 우편주소(등록 수리 서한)",
    ),
    48: (
        "E-mail address: 28576810@mylife.unisa.ac.za",
        "A5: myLife 이메일(사용자 제공)",
    ),
    49: (
        "Telephone nr(s): +82 10 7373 4960",
        "A5: 전화(사용자 제공)",
    ),
    59: (
        "(a) Initials & surname: E Mnkandla",
        "A6: Supervisor initials/surname(등록 서한). 서명은 교수용.",
    ),
    62: (
        "Telephone nr(s): (to be completed by supervisor)",
        "A6: 전화는 교수란 — 플레이스홀더만 표시",
    ),
    63: (
        "Address: School of Computing, College of Science, Engineering and Technology, "
        "University of South Africa",
        "A6: UNISA School of Computing",
    ),
    64: (
        "E-mail: mnkane@unisa.ac.za",
        "A6: 등록 서한 이메일",
    ),
    88: (
        "(a) Initials & surname: D K Moulla; D S Attipoe (additional co-supervisor)",
        "A8: 기존 co-supervisor Moulla 유지 + Attipoe 추가(교수 요청)",
    ),
    91: (
        "Institution: University of South Africa (School of Computing)",
        "A8: Institution = University of South Africa(사용자 확인)",
    ),
    92: (
        "Telephone nr(s) / Cell/Mobile: Attipoe office 011 471 2418 "
        "(Moulla: to be confirmed)",
        "A8: Attipoe 연락처(SoC staff page); Moulla 전화은 미확인",
    ),
    93: (
        "Address: Unisa Science Campus, Florida, Johannesburg "
        "(Attipoe: GJ Gerwel building, C3-028)",
        "A8: Attipoe 캠퍼스/사무실(SoC staff page)",
    ),
    95: (
        "Postal code: (UNISA Science Campus, Florida)",
        "A8: 우편번호 상세 미확인 — 캠퍼스만 표기",
    ),
    96: (
        "E-mail: moulladonatien@gmail.com; attipds@unisa.ac.za",
        "A8: Moulla(등록 서한) + Attipoe(attipds@unisa.ac.za)",
    ),
    121: (
        "(Not applicable — distance student in the Republic of Korea)",
        "거리 학생: Gauteng 인근 조항 비적용",
    ),
    122: ("N/A", "거리 학생: 월간 대면 조항 비적용"),
    123: ("N/A", "거리 학생: 인근 학생 통신 조항 비적용"),
    124: ("N/A", "거리 학생: 인근 학생 전화 조항 비적용"),
    125: (
        "(Applicable — distance student)",
        "거리 학생: 원격 조항 적용 표시",
    ),
    164: (
        "None / not applicable (no approved financial assistance specified for this study).",
        "B3: 펀딩 해당 없음",
    ),
    172: (
        B41,
        "B4.1: 후보 기대사항(거리 학생·컴퓨팅 시설 접근)",
    ),
    186: (
        B51,
        "B5.1: 연구 계획·시간 투입(제안서·논문·투고·심사)",
    ),
    195: (
        "None beyond the standard Procedures for Master's and Doctoral Students.",
        "IP: 특약 없음",
    ),
}


def set_paragraph_text(paragraph, text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def ensure_comments_part(docx_path: Path) -> None:
    tmp = docx_path.with_suffix(".comments_fix.docx")
    with zipfile.ZipFile(docx_path, "r") as zin, zipfile.ZipFile(
        tmp, "w", compression=zipfile.ZIP_DEFLATED
    ) as zout:
        names = set(zin.namelist())
        ct = etree.fromstring(zin.read("[Content_Types].xml"))
        rels = etree.fromstring(zin.read("word/_rels/document.xml.rels"))

        if not any(el.get("PartName") == "/word/comments.xml" for el in ct):
            etree.SubElement(
                ct,
                "{%s}Override" % CT_NS,
                PartName="/word/comments.xml",
                ContentType=(
                    "application/vnd.openxmlformats-officedocument"
                    ".wordprocessingml.comments+xml"
                ),
            )

        if not any(el.get("Type") == REL_COMMENTS for el in rels):
            used = {el.get("Id") for el in rels}
            rid = "rIdComments"
            n = 90
            while rid in used:
                rid = f"rId{n}"
                n += 1
            etree.SubElement(
                rels,
                "{%s}Relationship" % REL_NS,
                Id=rid,
                Type=REL_COMMENTS,
                Target="comments.xml",
            )

        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = etree.tostring(
                    ct, xml_declaration=True, encoding="UTF-8", standalone=True
                )
            elif item.filename == "word/_rels/document.xml.rels":
                data = etree.tostring(
                    rels, xml_declaration=True, encoding="UTF-8", standalone=True
                )
            elif item.filename == "word/comments.xml":
                continue
            zout.writestr(item, data)

        if "word/comments.xml" not in names:
            comments = etree.Element("{%s}comments" % W_NS, nsmap={"w": W_NS})
            zout.writestr(
                "word/comments.xml",
                etree.tostring(
                    comments, xml_declaration=True, encoding="UTF-8", standalone=True
                ),
            )

    shutil.move(str(tmp), str(docx_path))


def next_comment_id(comments_root) -> int:
    ids = [
        int(c.get(qn("w:id")))
        for c in comments_root.findall(qn("w:comment"))
        if c.get(qn("w:id")) is not None
    ]
    return (max(ids) + 1) if ids else 0


def add_comment_element(comments_root, cid: int, text: str) -> None:
    c = OxmlElement("w:comment")
    c.set(qn("w:id"), str(cid))
    c.set(qn("w:author"), AUTHOR)
    c.set(qn("w:date"), COMMENT_DATE)
    c.set(qn("w:initials"), INITIALS)
    p = OxmlElement("w:p")
    r0 = OxmlElement("w:r")
    r0.append(OxmlElement("w:annotationRef"))
    p.append(r0)
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    if text.startswith(" ") or text.endswith(" ") or "  " in text:
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = text
    r.append(t)
    p.append(r)
    c.append(p)
    comments_root.append(c)


def wrap_paragraph_comment(paragraph, cid: int) -> None:
    p = paragraph._p
    if p.find(qn("w:commentRangeStart")) is not None:
        return
    start = OxmlElement("w:commentRangeStart")
    start.set(qn("w:id"), str(cid))
    end = OxmlElement("w:commentRangeEnd")
    end.set(qn("w:id"), str(cid))
    ref_r = OxmlElement("w:r")
    ref = OxmlElement("w:commentReference")
    ref.set(qn("w:id"), str(cid))
    ref_r.append(ref)
    p.insert(0, start)
    p.append(end)
    p.append(ref_r)


def main() -> None:
    if not SRC.exists():
        raise SystemExit(f"Missing source: {SRC}")

    # Prefer blank backup if we already filled once
    if BACKUP.exists():
        shutil.copy2(BACKUP, SRC)
    else:
        shutil.copy2(SRC, BACKUP)

    ensure_comments_part(SRC)
    doc = Document(str(SRC))

    comments_part = None
    for rel in doc.part.rels.values():
        if rel.reltype == REL_COMMENTS:
            comments_part = rel.target_part
            break
    if comments_part is None:
        raise RuntimeError("comments part missing after ensure_comments_part")

    comments_root = comments_part.element
    cid = next_comment_id(comments_root)
    changed = 0

    for idx, (text, comment) in FILLS.items():
        if idx >= len(doc.paragraphs):
            raise SystemExit(f"Paragraph index out of range: {idx}")
        set_paragraph_text(doc.paragraphs[idx], text)
        add_comment_element(comments_root, cid, comment)
        wrap_paragraph_comment(doc.paragraphs[idx], cid)
        cid += 1
        changed += 1

    try:
        doc.save(str(OUT))
        out_used = OUT
    except PermissionError:
        alt = OUT.with_name(OUT.stem + "_filled.docx")
        doc.save(str(alt))
        out_used = alt
        print(f"PermissionError on {OUT}; saved to {alt}")

    REPO_COPY.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(out_used, REPO_COPY)

    with zipfile.ZipFile(out_used) as z:
        assert z.testzip() is None
        assert "word/comments.xml" in z.namelist()
        root = etree.fromstring(z.read("word/comments.xml"))
        authors = {c.get(qn("w:author")) for c in root.findall(qn("w:comment"))}
        n_comments = len(root.findall(qn("w:comment")))
        body = z.read("word/document.xml").decode("utf-8")
        assert "ns0:" not in body
        for s in [
            "Taehong Kwon",
            "28576810",
            "28576810@mylife.unisa.ac.za",
            "+82 10 7373 4960",
            "Attipoe",
            "TFCOS01",
            "mnkane@unisa.ac.za",
            "attipds@unisa.ac.za",
        ]:
            assert s in body, f"missing {s}"
        assert authors == {AUTHOR}
        print(f"OK comments={n_comments} authors={authors} changed={changed}")

    doc2 = Document(str(out_used))
    assert "Taehong Kwon" in doc2.paragraphs[9].text
    assert "28576810" in doc2.paragraphs[45].text
    assert "Attipoe" in doc2.paragraphs[88].text
    assert "distance student" in doc2.paragraphs[172].text.lower()
    print(f"Wrote {out_used}")
    print(f"Backup {BACKUP}")
    print(f"Copy {REPO_COPY}")


if __name__ == "__main__":
    main()
