# -*- coding: utf-8 -*-
"""Polish cover/TOC/short labels in the KO proposal DOCX (never touch EN source)."""
from __future__ import annotations

import re
import sys
import zipfile
from pathlib import Path

from deep_translator import GoogleTranslator
from docx import Document

SRC = Path(r"c:\Users\abito\Downloads\Proposal Final_080426_1705.docx")
DST = Path(r"c:\Users\abito\Downloads\Proposal Final_080426_1705_KO.docx")
CACHE = Path(__file__).resolve().parent / "_ko_translation_cache.json"

MANUAL = {
    "by": "지은이",
    "TAEHONG KWON": "권태홍 (TAEHONG KWON)",
    "submitted in accordance with the requirements for the degree of": "다음 학위 요건에 따라 제출함",
    "DOCTOR OF PHILOSOPHY": "철학박사 (Doctor of Philosophy)",
    "in the subject": "전공",
    "COMPUTER SCIENCE": "컴퓨터과학 (Computer Science)",
    "at the": "",
    "UNIVERSITY OF SOUTH AFRICA": "남아프리카공화국 대학교 (University of South Africa)",
    "SUPERVISOR: Prof. Ernest Mnkandla": "지도교수: Ernest Mnkandla 교수",
    "CO-SUPERVISOR: Prof. Donatien Koulla Moulla": "공동지도교수: Donatien Koulla Moulla 교수",
    "August 2026": "2026년 8월",
    "Abstract": "초록",
    "Table of Contents": "목차",
    "Introduction and background": "서론 및 배경",
    "Research Problem": "연구 문제",
    "Research Questions": "연구 질문",
    "Objectives": "연구 목표",
    "Specifications:": "명세:",
    "Keywords— DeFi, On-Chain Reputation, GMX V2, Perpetual Swap Positions, Arbitrum One, Blockchain Data Analytics, Wallet Endorsement Graphs": "키워드— DeFi, 온체인 평판, GMX V2, 무기한 스왑 포지션, Arbitrum One, 블록체인 데이터 분석, 지갑 보증 그래프",
    "INTEGRATING ERC-20 ALLOWANCE EDGES INTO PAGERANK FOR ENHANCED ON-CHAIN WALLET REPUTATION SCORING": "향상된 온체인 지갑 평판 산정을 위한 ERC-20 승인(allowance) 간선과 PageRank의 통합",
}


def set_text(paragraph, text: str) -> None:
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(text)
        return
    runs[0].text = text
    for r in runs[1:]:
        r.text = ""


def translate_left(translator: GoogleTranslator, left: str) -> str:
    left = left.strip()
    if not left:
        return left
    if left in MANUAL:
        return MANUAL[left]
    try:
        return translator.translate(left)
    except Exception:
        return left


def main() -> int:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    src = Document(str(SRC))
    ko = Document(str(DST))
    if len(src.paragraphs) != len(ko.paragraphs):
        print("paragraph count mismatch", len(src.paragraphs), len(ko.paragraphs))
        return 1

    translator = GoogleTranslator(source="en", target="ko")
    fixed = 0

    for sp, kp in zip(src.paragraphs, ko.paragraphs):
        st = sp.text
        if not st.strip():
            continue

        # Exact manual overrides
        if st.strip() in MANUAL:
            set_text(kp, MANUAL[st.strip()])
            fixed += 1
            continue

        # TOC / tab-separated label + page (or SPICE rows)
        if "\t" in st:
            left, right = st.split("\t", 1)
            left_ko = translate_left(translator, left)
            # keep page / right side as in source (numbers)
            set_text(kp, f"{left_ko}\t{right}")
            fixed += 1
            continue

    # Drop empty "at the" line content already set to ""
    ko.save(str(DST))

    with zipfile.ZipFile(DST, "r") as z:
        assert z.testzip() is None

    d2 = Document(str(DST))
    print("fixed paragraphs:", fixed)
    for i, p in enumerate(d2.paragraphs):
        if i in (2, 4, 7, 10, 13, 15, 18, 21, 24, 27, 29, 33, 35, 49, 51, 52, 54):
            print(i, repr(p.text))
    print("OK", DST)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
