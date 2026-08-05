# -*- coding: utf-8 -*-
"""Build bilingual EN+KO proposal DOCX: each English paragraph followed by its Korean.

Matching is by exact English paragraph text → translation map (cache + manual),
never by sequential index into a separate KO file.
"""
from __future__ import annotations

import json
import re
import shutil
import sys
import time
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from deep_translator import GoogleTranslator
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from lxml import etree

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

SRC = Path(r"c:\Users\abito\Downloads\Proposal Final_080426_1705.docx")
DST = Path(r"c:\Users\abito\Downloads\Proposal Final_080426_1705_EN-KO.docx")
CACHE = Path(__file__).resolve().parent / "_ko_translation_cache.json"

PROTECT = [
    "Adaptive Weighted PageRank",
    "EndorseRank",
    "PageRank",
    "RiskProp",
    "Arbitrum One",
    "BigQuery",
    "PositionDecrease",
    "PositionIncrease",
    "Human Passport",
    "Gitcoin Passport",
    "basePnlUsd",
    "orderType",
    "ERC-20",
    "EIP-2612",
    "GMX V2",
    "GMX",
    "AWP",
    "DeFi",
    "Sybil",
    "SPICE",
    "UNISA",
    "CSET",
    "Kendall",
    "Spearman",
    "Taehong Kwon",
    "Ernest Mnkandla",
    "Donatien Koulla Moulla",
    "David Sena Attipoe",
]

# High-visibility short labels (exact EN → KO)
MANUAL = {
    "by": "지은이",
    "TAEHONG KWON": "권태홍 (TAEHONG KWON)",
    "submitted in accordance with the requirements for the degree of": "다음 학위 요건에 따라 제출함",
    "DOCTOR OF PHILOSOPHY": "철학박사 (Doctor of Philosophy)",
    "in the subject": "전공",
    "COMPUTER SCIENCE": "컴퓨터과학 (Computer Science)",
    "at the": "（기관）",
    "UNIVERSITY OF SOUTH AFRICA": "남아프리카공화국 대학교 (University of South Africa)",
    "SUPERVISOR: Prof. Ernest Mnkandla": "지도교수: Ernest Mnkandla 교수",
    "CO-SUPERVISOR: Prof. Donatien Koulla Moulla": "공동지도교수: Donatien Koulla Moulla 교수",
    "August 2026": "2026년 8월",
    "Abstract": "초록",
    "Table of Contents": "목차",
    "Introduction and background": "서론 및 배경",
    "Research Problem": "연구 문제",
    "Research Questions": "연구 질문",
    "Objectives": "연구 목표",
    "Specifications:": "명세:",
    "Data Analysis\t51": "데이터 분석\t51",
    "Appendices\t55": "부록\t55",
    "INTEGRATING ERC-20 ALLOWANCE EDGES INTO PAGERANK FOR ENHANCED ON-CHAIN WALLET REPUTATION SCORING": (
        "향상된 온체인 지갑 평판 산정을 위한 ERC-20 승인(allowance) 간선과 PageRank의 통합"
    ),
    "Keywords— DeFi, On-Chain Reputation, GMX V2, Perpetual Swap Positions, Arbitrum One, Blockchain Data Analytics, Wallet Endorsement Graph": (
        "키워드— DeFi, 온체인 평판, GMX V2, 무기한 스왑 포지션, Arbitrum One, 블록체인 데이터 분석, 지갑 보증 그래프"
    ),
}


def _load_cache() -> dict[str, str]:
    if CACHE.exists():
        return json.loads(CACHE.read_text(encoding="utf-8"))
    return {}


def _save_cache(cache: dict[str, str]) -> None:
    CACHE.write_text(json.dumps(cache, ensure_ascii=False, indent=1), encoding="utf-8")


def _should_skip_translation(text: str) -> bool:
    t = text.strip()
    if not t:
        return True
    if re.fullmatch(r"[\d\W_]+", t, flags=re.UNICODE):
        return True
    return False


def _protect(text: str) -> tuple[str, list[str]]:
    held: list[str] = []
    out = text
    for term in PROTECT:
        while term in out:
            tok = f"⟦T{len(held)}⟧"
            held.append(term)
            out = out.replace(term, tok, 1)
    return out, held


def _unprotect(text: str, held: list[str]) -> str:
    out = text
    for i, val in enumerate(held):
        for tok in (f"⟦T{i}⟧", f"[T{i}]", f"(T{i})"):
            out = out.replace(tok, val)
    return out


def _translate_fresh(translator: GoogleTranslator, text: str) -> str:
    protected, held = _protect(text)
    if "\t" in protected:
        parts = protected.split("\t")
        out_parts = []
        for part in parts:
            if _should_skip_translation(part) and re.search(r"\d", part):
                out_parts.append(part)
            elif not part.strip():
                out_parts.append(part)
            else:
                out_parts.append(translator.translate(part))
                time.sleep(0.1)
        return _unprotect("\t".join(out_parts), held)
    return _unprotect(translator.translate(protected), held)


def resolve_ko(en: str, cache: dict[str, str], translator: GoogleTranslator) -> str | None:
    key = en
    stripped = en.strip()
    if stripped in MANUAL:
        return MANUAL[stripped]
    if key in MANUAL:
        return MANUAL[key]
    if key in cache:
        return cache[key]
    if _should_skip_translation(en):
        return None
    ko = _translate_fresh(translator, en)
    cache[key] = ko
    return ko


def insert_paragraph_after(paragraph: Paragraph, text: str, *, italic: bool = True) -> Paragraph:
    """Insert a new paragraph immediately after `paragraph`."""
    new_p = OxmlElement("w:p")
    # clone paragraph properties (style) when present
    pPr = paragraph._p.find(qn("w:pPr"))
    if pPr is not None:
        new_p.insert(0, etree.fromstring(etree.tostring(pPr)))
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    run = new_para.add_run(text)
    run.italic = italic
    return new_para


def append_title_comment(dst: Path) -> None:
    with zipfile.ZipFile(dst, "r") as zin:
        blob = {n: zin.read(n) for n in zin.namelist()}
    if "word/comments.xml" not in blob:
        return
    comments = etree.fromstring(blob["word/comments.xml"])
    ids = [int(c.get(f"{W}id")) for c in comments.findall(f"{W}comment") if c.get(f"{W}id")]
    new_id = str(max(ids) + 1 if ids else 0)
    now = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    note = (
        "요청: 원본 각 영문 문단 바로 아래에 대응 한국어 번역을 배치한 EN-KO 대역본 생성. "
        "문단 텍스트 키로 1:1 매칭(순번 혼선 방지). 원본/기존 KO 전용본은 덮어쓰지 않음."
    )
    comment = OxmlElement("w:comment")
    comment.set(qn("w:id"), new_id)
    comment.set(qn("w:author"), "Taehong Kwon")
    comment.set(qn("w:initials"), "TK")
    comment.set(qn("w:date"), now)
    p = OxmlElement("w:p")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = note
    r.append(t)
    p.append(r)
    comment.append(p)
    comments.append(comment)
    blob["word/comments.xml"] = etree.tostring(comments, xml_declaration=True, encoding="UTF-8", standalone=True)

    doc = etree.fromstring(blob["word/document.xml"])
    body = doc.find(f"{W}body")
    target = None
    for child in body:
        if etree.QName(child).localname != "p":
            continue
        if child.find(f".//{W}t") is not None:
            target = child
            break
    if target is not None:
        cstart = OxmlElement("w:commentRangeStart")
        cstart.set(qn("w:id"), new_id)
        cend = OxmlElement("w:commentRangeEnd")
        cend.set(qn("w:id"), new_id)
        cref_r = OxmlElement("w:r")
        cref = OxmlElement("w:commentReference")
        cref.set(qn("w:id"), new_id)
        cref_r.append(cref)
        target.insert(0, cstart)
        target.append(cend)
        target.append(cref_r)
        blob["word/document.xml"] = etree.tostring(doc, xml_declaration=True, encoding="UTF-8", standalone=True)

    tmp = dst.with_suffix(".tmp.zip")
    with zipfile.ZipFile(tmp, "w", compression=zipfile.ZIP_DEFLATED) as zout:
        for n, data in blob.items():
            zout.writestr(n, data)
    tmp.replace(dst)


def main() -> int:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    if not SRC.exists():
        print("Missing source", SRC)
        return 1
    if DST.exists():
        DST.unlink()

    shutil.copy2(SRC, DST)
    cache = _load_cache()
    translator = GoogleTranslator(source="en", target="ko")

    doc = Document(str(DST))
    # Snapshot English paragraphs first (before insertion mutates the list)
    en_paras = list(doc.paragraphs)
    pairs: list[tuple[Paragraph, str, str]] = []
    missing = 0
    for p in en_paras:
        en = p.text
        if not en.strip():
            continue
        ko = resolve_ko(en, cache, translator)
        if ko is None:
            continue
        if ko.strip() == en.strip():
            # no useful translation
            continue
        pairs.append((p, en, ko))

    # Insert from bottom to top so earlier anchors stay valid
    for p, en, ko in reversed(pairs):
        insert_paragraph_after(p, ko, italic=True)

    _save_cache(cache)
    doc.save(str(DST))
    append_title_comment(DST)

    with zipfile.ZipFile(DST, "r") as z:
        bad = z.testzip()
        print("testzip:", bad)
        assert bad is None
        assert b"ns0:" not in z.read("word/document.xml")

    # Verify adjacent EN→KO matching on a sample of body paragraphs
    d2 = Document(str(DST))
    texts = [p.text for p in d2.paragraphs]
    checks = 0
    mismatches = 0
    i = 0
    while i < len(texts) - 1:
        en = texts[i]
        if not en.strip():
            i += 1
            continue
        # expected: next non-empty after EN in bilingual pairs is KO for that EN
        # After our insert, KO is immediately next paragraph (even if empty? no, we insert right after)
        nxt = texts[i + 1] if i + 1 < len(texts) else ""
        expected = resolve_ko(en, cache, translator)
        if expected and nxt == expected:
            checks += 1
            i += 2
            continue
        if expected and nxt.strip():
            # possible mismatch if original had consecutive EN without our KO
            # Only flag when we know we inserted for this EN
            if en.strip() and expected.strip() and nxt != expected:
                # If next looks like Korean and we expected KO, count mismatch
                hangul = len(re.findall(r"[\uac00-\ud7a3]", nxt))
                if hangul > 5 and nxt != expected:
                    mismatches += 1
                    if mismatches <= 5:
                        print("MISMATCH near:", repr(en[:60]))
                        print("  expected:", repr(expected[:60]))
                        print("  got:     ", repr(nxt[:60]))
        i += 1

    print(f"Inserted KO under {len(pairs)} English paragraphs")
    print(f"Adjacent match checks OK≈{checks}, flagged mismatches≈{mismatches}")
    print("Sample (first 12 non-empty):")
    n = 0
    for t in texts:
        if not t.strip():
            continue
        print(" ", ("KO" if re.search(r"[\uac00-\ud7a3]", t) else "EN"), t[:100])
        n += 1
        if n >= 12:
            break
    print("DONE:", DST)
    print("bytes:", DST.stat().st_size)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
